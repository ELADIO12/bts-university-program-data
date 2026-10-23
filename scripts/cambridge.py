"""Cambridge workbook intake and research-only static publishing.

Workbook values are reported observations, not verified university facts.
This module never fetches websites or contacts BTS.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import html
import io
import json
import re
import unicodedata
from datetime import date, datetime
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urlsplit

from jsonschema import Draft202012Validator, FormatChecker

from . import bts

ROOT = Path(__file__).resolve().parents[1]
DATA = Path("data/cambridge/undergraduate")
SHEET = "BTS Fields - 32 Courses"
WORKBOOK_NAME = "Cambridge_Undergraduate_BTS_Fields_2027_Research.xlsx"
COLLECTION = "programs/cambridge/undergraduate"
SOURCE_COLUMNS = ("AY", "AZ", "BA")
SOURCE_FIELD = {"AY": "Degree / Course Source", "AZ": "Tuition Source", "BA": "Admissions Source"}
DIRECT_SOURCE = "Program Link"
COURSE_INDEX = "https://www.undergraduate.study.cam.ac.uk/courses/a-z"
LANGUAGE_CLASSES = {"VERIFIED_100_PERCENT_ENGLISH", "PARTIALLY_ENGLISH", "NON_ENGLISH", "UNVERIFIED_LANGUAGE"}
FLAG_COLUMNS = {
    "entrance_exam": "Entrance Exam",
    "gmat_gre_required": "GMAT/GRE Required",
    "numerus_fixus": "Numerus Fixus",
    "pre_master": "Pre-Master",
    "foundation_year": "Foundation Year",
    "post_bachelor_work_experience": "Post-Bachelor Work Experience",
    "erasmus_mundus": "Erasmus Mundus",
    "american_school": "American School",
    "franchise": "Franchise",
    "sat_act_required": "SAT/ACT Required",
    "spring_intake": "Spring Intake",
    "internship_built_in": "Internship Built In",
}
DIRECT_COLUMNS = {
    "program_title": "Program Title",
    "degree_level": "Degree Level",
    "program_link": "Program Link",
    "contact_university": "Contact (University)",
    "degree_type": "Degree Type",
    "areas_of_study": "Areas of Study",
    "primary_area_of_study": "Primary Area of Study",
    "additional_areas_of_study": "Additional Areas of Study",
    "description_html": "Description (HTML)",
    "university_street": "Street",
    "university_street_2": "Street 2",
    "university_postal_code": "Postal Code",
    "university_city": "City",
    "university_region": "Region",
    "university_country": "Country",
    "university_latitude": "Latitude",
    "university_longitude": "Longitude",
    "application_deadline": "Application Deadline",
    "application_deadline_date": "Application Deadline Date",
    "start_application_url": "Start Your Application",
    "admissions_requirements": "Admissions Requirements",
    "admissions_difficulty": "Admissions Difficulty",
    "minimum_gpa": "Minimum GPA",
    "tuition_source_url": "Tuition Source",
}


def _slug(value: str) -> str:
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode().lower()
    return "-".join(re.findall(r"[a-z0-9]+", value))


def _json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _uri(value: object) -> bool:
    if not isinstance(value, str):
        return False
    parts = urlsplit(value)
    return parts.scheme == "https" and bool(parts.hostname)


def _official(value: object) -> bool:
    return _uri(value) and (urlsplit(value).hostname == "cam.ac.uk" or urlsplit(value).hostname.endswith(".cam.ac.uk"))


def extract_workbook(path: Path) -> dict:
    from openpyxl import load_workbook

    source = Path(path)
    workbook = load_workbook(source, data_only=False, read_only=False)
    sheets = {}
    for ws in workbook.worksheets:
        rows = []
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            cells = []
            for cell in row:
                if cell.data_type == "f":
                    raise ValueError(f"Unexpected formula in {ws.title}!{cell.coordinate}")
                value = cell.value
                if isinstance(value, (datetime, date)):
                    value = value.isoformat()
                cells.append(value)
            rows.append(cells)
        sheets[ws.title] = {"max_row": ws.max_row, "max_column": ws.max_column, "rows": rows}
    if set(sheets) != {SHEET, "READ ME", "Tuition bands 2027"}:
        raise ValueError("Unexpected workbook sheets")
    return {
        "filename": source.name,
        "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "sheets": sheets,
        "note": "Exact worksheet cell values as supplied. Dates/formulas are serialized without changing substantive content."
    }


def _row(raw: dict, excel_row: int) -> tuple[dict, dict]:
    sheet = raw["sheets"][SHEET]["rows"]
    headers = sheet[0]
    values = sheet[excel_row - 1]
    if len(headers) != 55:
        raise ValueError("Expected 55 workbook columns")
    by_name = dict(zip(headers, values))
    by_cell = {header: f"{_column(i + 1)}{excel_row}" for i, header in enumerate(headers)}
    return by_name, by_cell


def _column(n: int) -> str:
    letters = ""
    while n:
        n, rem = divmod(n - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def _field(value: object, cell: str | None, note: str = "") -> dict:
    if isinstance(value, str) and (not value.strip() or value.strip().lower() in {"unverified", "unknown", "n/a"}):
        value = None
    return {
        "value": value,
        "status": "workbook_reported" if value is not None else "unverified",
        "source_cell": cell if value is not None else None,
        "review_note": note or ("Reported by supplied workbook; independent official-source review pending." if value is not None else "No supported workbook value.")
    }


def _flag(value: object) -> bool | None:
    if isinstance(value, str):
        v = value.strip().lower()
        if v == "yes":
            return True
        if v == "no":
            return False
    return None


def _choice(row: dict, cells: dict, yes_columns: tuple[tuple[str, str], ...]) -> dict:
    selected = [(label, cells[col]) for col, label in yes_columns if _flag(row[col]) is True]
    if len(selected) == 1:
        return _field(selected[0][0], selected[0][1])
    return _field(None, None, "Multiple or no choices marked Yes; manual review required.")


def _fee(amount: object, cell: str, kind: str) -> dict:
    if not isinstance(amount, (int, float)) or isinstance(amount, bool):
        return _field(None, None, f"No supported {kind} annual tuition amount.")
    return _field({
        "amount": amount,
        "currency": "GBP",
        "period": "year",
        "entry_year": 2027,
        "academic_year": None,
        "fee_status": kind,
        "note": "Workbook amount only; annual University tuition. College fees excluded. Official 2027 source could not be accessed for independent verification."
    }, cell, "2027 entry is stated in the workbook; exact academic-year label requires source review.")


def _date(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, str):
        try:
            return datetime.strptime(value, "%d/%m/%Y").date().isoformat()
        except ValueError:
            return None
    return None


def convert(raw: dict) -> list[dict]:
    rows = raw["sheets"][SHEET]["rows"]
    if len(rows) != 33:
        raise ValueError(f"Expected 32 Cambridge course rows, found {len(rows) - 1}")
    records = []
    for excel_row in range(2, 34):
        row, cells = _row(raw, excel_row)
        title = row["Program Title"]
        qualification = row["Degree Type"]
        if not isinstance(title, str) or not title.strip() or not isinstance(qualification, str) or not qualification.strip():
            raise ValueError(f"Missing Cambridge identity in row {excel_row}")
        rid = f"cambridge-{_slug(title)}-{_slug(qualification)}"
        fields = {key: _field(None, None) for key in bts.LABELS}
        for key, col in DIRECT_COLUMNS.items():
            value = row[col]
            if key == "application_deadline_date":
                value = _date(value)
            if key == "degree_level" and value == "Bachelors":
                value = "Bachelors"
            fields[key] = _field(value, cells[col])
        # School Name contains the university name, not an identifiable school.
        fields["school_name"] = _field(None, None, "Workbook names the university, not a program school or faculty.")
        fields["format"] = _choice(row, cells, (("Full Time", "Full Time"), ("Part Time", "Part Time")))
        fields["attendance"] = _choice(row, cells, (("On Campus", "On Campus"), ("Online", "Online"), ("Blended Learning", "Blended Learning")))
        exact = row["Exact Duration (years)"]
        fields["duration_original"] = _field(
            {"amount": exact, "unit": "years"} if isinstance(exact, (int, float)) else None,
            cells["Exact Duration (years)"],
            "Workbook route duration; integrated-degree and professional routes require individual course review."
        )
        bts_duration = row["Duration (BTS)"] if row["Duration (BTS)"] in bts.DURATION_CHOICES else None
        fields["duration_bts"] = _field(bts_duration, cells["Duration (BTS)"], "Normalized from workbook duration; course route requires review.")
        fields["tuition_home"] = _fee(row["Tuition (Home/yr GBP)"], cells["Tuition (Home/yr GBP)"], "Home")
        fields["tuition_non_eu"] = _fee(row["Tuition (non-EU/yr GBP)"], cells["Tuition (non-EU/yr GBP)"], "Overseas")
        fields["tuition_currency"] = _field("GBP", cells["Tuition (non-EU/yr GBP)"], "Currency is specified in workbook column heading.")
        fields["tuition_academic_year"] = _field(None, None, "Workbook states 2027 entry, but does not explicitly state an academic-year label.")
        fields["tuition_last_verified"] = _field(None, None, "No independent official-source verification date.")
        fields["tuition_billing_period"] = _field("year", cells["Tuition (non-EU/yr GBP)"], "Annual basis is specified in workbook column heading.")
        for key, col in FLAG_COLUMNS.items():
            fields[key] = _field(_flag(row[col]), cells[col], "Workbook assertion only; no official course-specific check.")
        fields["english_proficiency_required"] = _field(None, None, "Conditional workbook text does not establish a yes/no course-specific requirement.")
        sources = []
        for col in (DIRECT_SOURCE, *SOURCE_FIELD.values()):
            value = row[col]
            if isinstance(value, str) and _uri(value) and value not in sources:
                sources.append(value)
        direct = row[DIRECT_SOURCE] if _official(row[DIRECT_SOURCE]) else None
        course_source = row["Degree / Course Source"] if _official(row["Degree / Course Source"]) else COURSE_INDEX
        flags = ["Compulsory teaching language is not independently verified.",
                 "Three existing Cambridge BTS program identities were not supplied; duplicate check pending.",
                 "Official Cambridge pages returned HTTP 403 in this research session; manual source evidence is required."]
        if direct is None:
            flags.append("No verified direct course link in the workbook.")
        if row["Tuition (non-EU/yr GBP)"] is None:
            flags.append("Overseas tuition missing from workbook.")
        if exact == 4 and qualification == "BA (Hons)":
            flags.append("Four-year BA or integrated degree route needs course-specific qualification review.")
        if exact and exact > 4:
            flags.append("Professional degree length and year-specific fees need course review.")
        if title == "Medicine (Graduate course)":
            flags.append("Graduate-entry Medicine has special fee and application arrangements; workbook home fee and standard deadline require separate verification.")
        if str(row["BTS Import Status"]).startswith("EXCLUDE"):
            flags.append("Supplied workbook explicitly excludes this course from BTS import.")
        if row["Application Deadline Date"] == "15/10/2026":
            flags.append("Workbook standard 2027-entry deadline is 15 October 2026; check exceptions and any later intake.")
        if row["Notes / Review Flags"]:
            flags.append(str(row["Notes / Review Flags"]))
        record = {
            "schema_version": "cambridge-research-1.0",
            "id": rid,
            "university": {
                "name": "University of Cambridge",
                "identity_domain": "cam.ac.uk",
                "bts_university_id": 92862,
                "id_provenance": "user-provided; not checked against BTS"
            },
            "source_workbook": {
                "filename": raw["filename"],
                "sha256": raw["sha256"],
                "sheet": SHEET,
                "row": excel_row
            },
            "program": {
                "title": title,
                "qualification": qualification,
                "direct_program_url": direct,
                "research_source_url": course_source
            },
            "bts_fields": fields,
            "official_sources": sources,
            "language_classification": row["English Instruction Verification"] if row["English Instruction Verification"] in LANGUAGE_CLASSES else "UNVERIFIED_LANGUAGE",
            "existing_bts_match": "unknown",
            "workbook_import_status": str(row["BTS Import Status"]),
            "import_eligibility": "hold",
            "human_review_status": "pending",
            "review_flags": list(dict.fromkeys(flags))
        }
        records.append(record)
    return records


def summary(records: list[dict]) -> dict:
    return {
        "university_name": "University of Cambridge",
        "bts_university_id": 92862,
        "record_count": len(records),
        "verified_100_percent_english": sum(r["language_classification"] == "VERIFIED_100_PERCENT_ENGLISH" for r in records),
        "language_review_required": sum(r["language_classification"] == "UNVERIFIED_LANGUAGE" for r in records),
        "excluded_non_english_or_partial": sum(r["language_classification"] in {"NON_ENGLISH", "PARTIALLY_ENGLISH"} for r in records),
        "confirmed_existing_in_bts": sum(r["existing_bts_match"] == "confirmed_existing" for r in records),
        "workbook_explicit_exclusions": sum(r["workbook_import_status"].startswith("EXCLUDE") for r in records),
        "duplicate_check_pending": sum(r["existing_bts_match"] == "unknown" for r in records),
        "new_import_ready": sum(r["import_eligibility"] == "eligible_after_human_review" for r in records),
        "workbook_direct_course_links": sum(bool(r["program"]["direct_program_url"]) for r in records),
        "overseas_fee_missing_in_workbook": sum(r["bts_fields"]["tuition_non_eu"]["value"] is None for r in records),
        "official_source_access": "HTTP 403 in this research session; manual evidence collection required",
        "notice": "Counts describe the supplied workbook and research status, not BTS approval."
    }


def validation_report(records: list[dict]) -> dict:
    report = summary(records)
    report["checks"] = {
        "missing_titles": sum(not r["program"]["title"] for r in records),
        "duplicate_title_qualification": len(records) - len({(_slug(r["program"]["title"]), _slug(r["program"]["qualification"])) for r in records}),
        "missing_direct_course_links": sum(r["program"]["direct_program_url"] is None for r in records),
        "missing_overseas_tuition": report["overseas_fee_missing_in_workbook"],
        "unverified_teaching_language": report["language_review_required"],
        "existing_bts_comparison_pending": report["duplicate_check_pending"],
        "tuition_academic_year_unstated": sum(r["bts_fields"]["tuition_academic_year"]["value"] is None for r in records),
        "undergraduate_degree_classification_review": sum(r["program"]["title"] == "Medicine (Graduate course)" for r in records),
        "four_year_ba_route_review": sum(r["bts_fields"]["duration_original"]["value"]["amount"] == 4
                                         and r["program"]["qualification"] == "BA (Hons)" for r in records)
    }
    report["issues"] = [
        {"id": r["id"], "workbook_row": r["source_workbook"]["row"], "review_flags": r["review_flags"]}
        for r in records
    ]
    return report


def _csv(records: list[dict]) -> str:
    output = io.StringIO(newline="")
    columns = ["id", "title", "qualification", "workbook_row", "bts_university_id", "language_classification",
               "existing_bts_match", "import_eligibility", "direct_program_url"]
    for key in bts.LABELS:
        columns.extend((key, key + "_status", key + "_source_cell"))
    writer = csv.DictWriter(output, fieldnames=columns, lineterminator="\n")
    writer.writeheader()
    for r in records:
        row = {
            "id": r["id"], "title": r["program"]["title"], "qualification": r["program"]["qualification"],
            "workbook_row": r["source_workbook"]["row"], "bts_university_id": 92862,
            "language_classification": r["language_classification"],
            "existing_bts_match": r["existing_bts_match"],
            "import_eligibility": r["import_eligibility"],
            "direct_program_url": r["program"]["direct_program_url"] or ""
        }
        for key, field in r["bts_fields"].items():
            val = field["value"]
            text = "" if val is None else json.dumps(val, ensure_ascii=False) if isinstance(val, (dict, list)) else str(val)
            if text.lstrip().startswith(("=", "+", "-", "@")):
                text = "'" + text
            row[key] = text
            row[key + "_status"] = field["status"]
            row[key + "_source_cell"] = field["source_cell"] or ""
        writer.writerow(row)
    return output.getvalue()


def ingest(workbook_path: Path, root: Path = ROOT) -> list[dict]:
    raw = extract_workbook(workbook_path)
    records = convert(raw)
    target = root / DATA
    _json(target / "raw-workbook.json", raw)
    for r in records:
        _json(target / "programs" / (r["id"] + ".json"), r)
    _json(target / "cambridge-undergraduate.json", {
        "collection": "University of Cambridge undergraduate workbook research",
        "summary": summary(records),
        "records": records
    })
    read_me = raw["sheets"]["READ ME"]["rows"]
    collection_sources = [row[1] for row in read_me[11:16] if len(row) > 1 and _official(row[1])]
    _json(target / "discovery-manifest.json", {
        "research_directory": COLLECTION + "/",
        "import_ready_directory": COLLECTION + "/import-ready/",
        "single_page_extraction_test": COLLECTION + "/all-programs.html",
        "workbook_sha256": raw["sha256"],
        "collection_official_sources": collection_sources,
        "programs": [{"id": r["id"], "title": r["program"]["title"], "path": COLLECTION + "/" + filename(r)} for r in records]
    })
    _json(target / "excluded-programs.json", [
        {"id": r["id"], "title": r["program"]["title"],
         "status": "workbook_excluded" if r["workbook_import_status"].startswith("EXCLUDE") else "hold",
         "reasons": (["workbook explicitly excludes this course"] if r["workbook_import_status"].startswith("EXCLUDE") else [])
                    + ["language not certified 100% English", "existing BTS match not established"]}
        for r in records
    ])
    _json(target / "validation-report.json", validation_report(records))
    (target / "cambridge-undergraduate.csv").write_text(_csv(records), encoding="utf-8", newline="\n")
    validate(root)
    return records


def filename(record: dict) -> str:
    return record["id"].removeprefix("cambridge-") + ".html"


def load(root: Path = ROOT) -> tuple[list[dict], dict]:
    target = root / DATA
    if not target.exists():
        return [], {}
    raw = json.loads((target / "raw-workbook.json").read_text(encoding="utf-8"))
    records = [json.loads(p.read_text(encoding="utf-8")) for p in sorted((target / "programs").glob("*.json"))]
    return records, raw


def validate(root: Path = ROOT) -> dict:
    records, raw = load(root)
    if not records and not raw:
        return {}
    schema = json.loads((root / "schema/cambridge-program.schema.json").read_text(encoding="utf-8"))
    Draft202012Validator.check_schema(schema)
    validator = Draft202012Validator(schema, format_checker=FormatChecker())
    if len(records) != 32 or len(raw["sheets"][SHEET]["rows"]) != 33:
        raise ValueError("Cambridge input must contain exactly 32 course records")
    ids, rows, identities, direct_urls = set(), set(), set(), set()
    for record in records:
        validator.validate(record)
        rid = record["id"]
        if rid in ids:
            raise ValueError(f"Duplicate Cambridge id: {rid}")
        ids.add(rid)
        if set(record["bts_fields"]) != set(bts.LABELS):
            raise ValueError(f"{rid}: BTS 64-field coverage mismatch")
        source_row = record["source_workbook"]["row"]
        if source_row in rows:
            raise ValueError(f"Duplicate Cambridge workbook row: {source_row}")
        rows.add(source_row)
        input_row, cells = _row(raw, source_row)
        if record["source_workbook"]["sha256"] != raw["sha256"] or record["program"]["title"] != input_row["Program Title"] or record["program"]["qualification"] != input_row["Degree Type"]:
            raise ValueError(f"{rid}: workbook identity mismatch")
        if str(input_row["University ID"]).strip() != "92862":
            raise ValueError(f"{rid}: workbook university ID mismatch")
        if record["workbook_import_status"] != str(input_row["BTS Import Status"]):
            raise ValueError(f"{rid}: workbook import status mismatch")
        if record["program"]["direct_program_url"] != (input_row["Program Link"] if _official(input_row["Program Link"]) else None):
            raise ValueError(f"{rid}: direct program link mismatch")
        if rid != f"cambridge-{_slug(record['program']['title'])}-{_slug(record['program']['qualification'])}":
            raise ValueError(f"{rid}: unstable identifier")
        identity = (_slug(record["program"]["title"]), _slug(record["program"]["qualification"]))
        if identity in identities:
            raise ValueError(f"{rid}: duplicate title and qualification")
        identities.add(identity)
        for url in record["official_sources"]:
            if not _official(url):
                raise ValueError(f"{rid}: non-Cambridge official source")
        for key, field in record["bts_fields"].items():
            if field["status"] == "unverified" and field["value"] is not None:
                raise ValueError(f"{rid}.{key}: unverified value must be null")
            if field["status"] == "verified":
                raise ValueError(f"{rid}.{key}: independent evidence registry required before verified status")
            if field["source_cell"] and not field["source_cell"].endswith(str(source_row)):
                raise ValueError(f"{rid}.{key}: wrong source row")
        if record["bts_fields"]["tuition_non_eu"]["value"] is not None:
            if record["bts_fields"]["tuition_non_eu"]["value"]["amount"] != input_row["Tuition (non-EU/yr GBP)"]:
                raise ValueError(f"{rid}: overseas tuition differs from workbook")
        if record["bts_fields"]["duration_original"]["value"] is not None:
            if record["bts_fields"]["duration_original"]["value"]["amount"] != input_row["Exact Duration (years)"]:
                raise ValueError(f"{rid}: duration differs from workbook")
        if record["language_classification"] == "VERIFIED_100_PERCENT_ENGLISH" and record["existing_bts_match"] == "confirmed_new" and record["human_review_status"] == "reviewed":
            pass
        elif record["import_eligibility"] != "hold":
            raise ValueError(f"{rid}: unsupported import eligibility")
        if record["program"]["direct_program_url"] and not _official(record["program"]["direct_program_url"]):
            raise ValueError(f"{rid}: invalid Cambridge course link")
        if record["program"]["direct_program_url"]:
            canonical_url = record["program"]["direct_program_url"].rstrip("/").lower()
            if canonical_url in direct_urls:
                raise ValueError(f"{rid}: duplicate direct Cambridge course URL")
            direct_urls.add(canonical_url)
        if record["bts_fields"]["duration_bts"]["value"] != input_row["Duration (BTS)"]:
            raise ValueError(f"{rid}: BTS duration differs from workbook")
    if rows != set(range(2, 34)):
        raise ValueError("Cambridge workbook row coverage mismatch")
    return summary(records)


class _Text(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []

    def handle_data(self, data):
        self.parts.append(data)


def _display(value: object, key: str) -> str:
    if value is None:
        return "No supported value"
    if key == "description_html":
        parser = _Text()
        parser.feed(str(value))
        return html.escape(" ".join(" ".join(parser.parts).split()))
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, dict) and "amount" in value and "currency" in value:
        return html.escape(f"GBP {value['amount']:,} per year, 2027 entry; academic-year label unverified; College fees excluded")
    if isinstance(value, dict) and "unit" in value:
        return html.escape(f"{value['amount']} {value['unit']}")
    return html.escape(str(value))


def _document(title: str, body: str, stylesheet: str, canonical: str | None = None, course: dict | None = None, noindex: bool = False) -> str:
    extra = '<meta name="robots" content="noindex,nofollow">' if noindex else ""
    if canonical:
        extra += f'<link rel="canonical" href="{html.escape(canonical, quote=True)}">'
    if course:
        safe = json.dumps(course, ensure_ascii=False).replace("<", "\\u003c").replace(">", "\\u003e").replace("&", "\\u0026")
        extra += f'<script type="application/ld+json">{safe}</script>'
    return (f'<!doctype html><html lang="en"><head><meta charset="utf-8">'
            f'<meta name="viewport" content="width=device-width, initial-scale=1">'
            f'<title>{html.escape(title)} | Cambridge BTS research</title>'
            f'<meta name="description" content="Independent Cambridge undergraduate workbook research; not verified for BTS import.">'
            f'<link rel="stylesheet" href="{stylesheet}">{extra}</head><body>'
            f'<header><strong>BTS · Cambridge undergraduate research</strong><p>RESEARCH ONLY / EXPERIMENTAL</p></header>'
            f'<main><aside class="notice">Independent research material. This is not an official University of Cambridge website. '
            f'Workbook values have not been independently verified and no program is approved for BTS import.</aside>{body}</main>'
            f'<footer>Consult official Cambridge sources and a human reviewer before using any value.</footer></body></html>')


def _program_html(record: dict, base_url: str) -> str:
    esc = html.escape
    title = record["program"]["title"]
    degree = record["program"]["qualification"]
    source = record["program"]["direct_program_url"] or record["program"]["research_source_url"]
    source_label = "Direct course page listed in workbook" if record["program"]["direct_program_url"] else "Cambridge course directory; direct course link missing"
    sections = []
    for section, labels in bts.SECTIONS.items():
        items = []
        for key, label in labels.items():
            field = record["bts_fields"][key]
            value = _display(field["value"], key)
            if key in bts.URL_FIELDS and field["value"] is not None and _uri(field["value"]):
                value = f'<a href="{esc(field["value"], quote=True)}">{value}</a>'
            items.append(f'<dt>{esc(label)}</dt><dd>{value}<br><small>Status: {esc(field["status"])}; '
                         f'{esc(field["review_note"])}'
                         + (f' Source: workbook cell {esc(field["source_cell"])}.' if field["source_cell"] else "") + '</small></dd>')
        sections.append(f'<section><h2>{esc(section)}</h2><dl>{"".join(items)}</dl></section>')
    official = "".join(f'<li><a href="{esc(url, quote=True)}">{esc(url)}</a></li>' for url in record["official_sources"])
    flags = "".join(f'<li>{esc(flag)}</li>' for flag in record["review_flags"])
    body = (f'<h1>{esc(title)} — {esc(degree)}</h1>'
            f'<p>University of Cambridge · Undergraduate · BTS university ID 92862 (user-provided, not independently checked).</p>'
            f'<p><strong>Research status:</strong> language {esc(record["language_classification"])}; '
            f'BTS duplicate check pending; import eligibility on hold; human review pending. '
            f'Workbook import status: {esc(record["workbook_import_status"])}.</p>'
            f'<p><a href="{esc(source, quote=True)}">Original university source: {esc(source_label)}</a>. '
            f'Official page access returned HTTP 403 during this research pass; manual evidence collection is required.</p>'
            f'<p><a href="../../../data/cambridge/undergraduate/programs/{esc(record["id"])}.json">Structured research record (JSON)</a></p>'
            + "".join(sections)
            + f'<section><h2>Official source references</h2><ul>{official}</ul></section>'
            + f'<section><h2>Human review flags</h2><ul>{flags}</ul></section>')
    canonical = base_url.rstrip("/") + "/" + COLLECTION + "/" + filename(record) if base_url else None
    course = {
        "@context": "https://schema.org",
        "@type": "Course",
        "name": title,
        "identifier": record["id"],
        "provider": {"@type": "CollegeOrUniversity", "name": "University of Cambridge"},
        "isBasedOn": source,
        "additionalProperty": [
            {"@type": "PropertyValue", "name": "Research status", "value": "Workbook-reported, independent verification pending"},
            {"@type": "PropertyValue", "name": "BTS import eligibility", "value": "hold"}
        ]
    }
    if canonical:
        course["url"] = canonical
    return _document(title, body, "../../../assets/style.css", canonical, course, noindex=True)


def _all_programs_html(records: list[dict], base_url: str) -> str:
    """One self-contained page for manual, review-only importer compatibility testing."""
    esc = html.escape
    articles = []
    graph = []
    canonical = base_url.rstrip("/") + "/" + COLLECTION + "/all-programs.html" if base_url else ""
    compact_fields = (
        "areas_of_study", "primary_area_of_study", "description_html",
        "format", "attendance", "duration_original", "tuition_home",
        "tuition_non_eu", "application_deadline_date", "start_application_url",
        "admissions_requirements", "entrance_exam"
    )
    for record in records:
        program = record["program"]
        source = program["direct_program_url"] or program["research_source_url"]
        fields = []
        for key in compact_fields:
            label = bts.LABELS[key]
            field = record["bts_fields"][key]
            if field["value"] is None:
                continue
            fields.append(
                f'<dt>{esc(label)}</dt><dd data-status="workbook_reported" '
                f'data-source-cell="{esc(field["source_cell"] or "", quote=True)}">'
                f'{_display(field["value"], key)}</dd>'
            )
        generic = ("Compulsory teaching language", "Three existing Cambridge BTS",
                   "Official Cambridge pages", "No verified direct course link",
                   "Workbook standard 2027-entry deadline")
        flags = "; ".join(esc(flag) for flag in record["review_flags"] if not flag.startswith(generic))
        articles.append(
            f'<article class="program" id="{esc(record["id"], quote=True)}">'
            f'<h2>{esc(program["title"])} — {esc(program["qualification"])}</h2>'
            f'<p>University: University of Cambridge. BTS university ID: 92862 '
            f'(user-provided, not independently checked). Degree level: Bachelors '
            f'(workbook-reported). Internal research ID: {esc(record["id"])}.</p>'
            f'<p><strong>Status:</strong> research only; {esc(record["language_classification"])}; '
            f'existing BTS match unknown; not eligible for import; human review pending. '
            f'Workbook import status: {esc(record["workbook_import_status"])}.</p>'
            f'<p>Original Cambridge source URL (plain text): {esc(source)}. '
            f'Cambridge source pages attempted in this research pass returned HTTP 403; '
            f'course-specific evidence requires permitted manual collection.</p>'
            + '<p>Every populated field below is workbook-reported and awaits official verification. '
            'The Excel cell is recorded in each field’s source-cell attribute.</p>'
            + f'<dl>{"".join(fields)}</dl>'
            + (f'<p><strong>Course-specific review notes:</strong> {flags}</p>' if flags else "")
            + '</article>'
        )
        course = {
            "@type": "Course",
            "@id": (canonical + "#" + record["id"]) if canonical else "#" + record["id"],
            "name": program["title"],
            "identifier": record["id"],
            "provider": {"@type": "CollegeOrUniversity", "name": "University of Cambridge"},
            "additionalProperty": [
                {"@type": "PropertyValue", "name": "Research status", "value": "workbook_reported; official verification pending"},
                {"@type": "PropertyValue", "name": "BTS import eligibility", "value": "hold"}
            ]
        }
        graph.append(course)
    structured = json.dumps({"@context": "https://schema.org", "@graph": graph}, ensure_ascii=False)
    structured = structured.replace("<", "\\u003c").replace(">", "\\u003e").replace("&", "\\u0026")
    style = (
        "body{font:16px/1.55 system-ui,sans-serif;color:#1c2933;max-width:1000px;margin:auto;padding:2rem}"
        "h1,h2,h3{line-height:1.25}h2{margin-top:0}.notice{background:#fff2d8;border:2px solid #a66b00;padding:1rem}"
        ".program{border-top:3px solid #42647b;margin-top:2.5rem;padding-top:1.5rem}"
        "section{margin:1rem 0}dt{font-weight:700;margin-top:.6rem}dd{margin:0 0 .7rem 0}"
        "small{display:block;color:#52606d}li{margin:.25rem 0}code{overflow-wrap:anywhere}"
    )
    return (
        '<!doctype html><html lang="en"><head><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width, initial-scale=1">'
        '<title>Cambridge undergraduate programs | one-page research extraction test</title>'
        '<meta name="description" content="32 Cambridge undergraduate workbook research records on one self-contained page.">'
        f'<style>{style}</style><script type="application/ld+json">{structured}</script></head><body>'
        '<main><h1>University of Cambridge undergraduate programs — research extraction test</h1>'
        '<aside class="notice"><strong>Independent research material.</strong> This is not an official University of Cambridge website. '
        'All 32 records are workbook-reported, not independently verified. None is eligible for BTS import. '
        'Use this page only to test extraction and stop at the importer review stage without confirming any imports.</aside>'
        '<p>This self-contained page includes every program below. It has no links to other program pages and no external stylesheet or script.</p>'
        + "".join(articles)
        + '</main></body></html>'
    )


def site_files(records: list[dict]) -> set[str]:
    prefix = COLLECTION + "/"
    data = "data/cambridge/undergraduate/"
    names = {prefix + "index.html", prefix + "import-ready/index.html", prefix + "all-programs.html",
             data + "raw-workbook.json", data + "cambridge-undergraduate.json",
             data + "cambridge-undergraduate.csv", data + "discovery-manifest.json",
             data + "excluded-programs.json", data + "validation-report.json",
             "reports/cambridge/verification-report.md", "reports/cambridge/duplicate-check.md",
             "reports/cambridge/source-coverage.md"}
    for r in records:
        names.add(prefix + filename(r))
        names.add(data + "programs/" + r["id"] + ".json")
    return names


def build(root: Path, out: Path, base_url: str = "") -> set[str]:
    records, _ = load(root)
    if not records:
        return set()
    report = validate(root)
    expected = site_files(records)
    def write(relative: str, content: str) -> None:
        path = out / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(content, encoding="utf-8", newline="\n")
    links = []
    for r in records:
        write(COLLECTION + "/" + filename(r), _program_html(r, base_url))
        links.append(f'<li><a href="{html.escape(filename(r), quote=True)}">{html.escape(r["program"]["title"])}'
                     f' — {html.escape(r["program"]["qualification"])}</a>. '
                     f'Language: {r["language_classification"]}; import: hold.</li>')
    write(COLLECTION + "/all-programs.html", _all_programs_html(records, base_url))
    research_body = (
        '<h1>University of Cambridge — Undergraduate Degree Programs</h1>'
        f'<p>{len(records)} workbook course records. No course has independently verified 100% English teaching. '
        'The three existing Cambridge BTS program identities are unavailable. All records remain research-only.</p>'
        '<p><strong>Do not submit this complete research directory to the BTS importer.</strong></p>'
        '<p><a href="all-programs.html">One-page research extraction test</a> (all 32 records inline; review only).</p>'
        '<p><a href="import-ready/index.html">Separate BTS import-ready directory</a> (currently empty).</p>'
        '<p><a href="../../../data/cambridge/undergraduate/cambridge-undergraduate.csv">Download research CSV</a> · '
        '<a href="../../../data/cambridge/undergraduate/validation-report.json">Validation report</a></p>'
        '<h2>Course research pages</h2><ul class="programs">' + "".join(links) + '</ul>'
    )
    canonical = base_url.rstrip("/") + "/" + COLLECTION + "/" if base_url else None
    write(COLLECTION + "/index.html", _document(
        "University of Cambridge — Undergraduate Degree Programs", research_body,
        "../../../assets/style.css", canonical, noindex=True))
    ready_body = ('<h1>Cambridge undergraduate programs eligible for BTS review</h1>'
                  '<p>No programs currently qualify. All 32 workbook records await compulsory teaching-language '
                  'verification and comparison with the three existing Cambridge BTS records.</p>'
                  '<p>There are no program links on this page. Do not initiate an import from this empty directory.</p>')
    ready_canonical = base_url.rstrip("/") + "/" + COLLECTION + "/import-ready/" if base_url else None
    write(COLLECTION + "/import-ready/index.html", _document(
        "Cambridge BTS import-ready directory", ready_body, "../../../../assets/style.css", ready_canonical))
    for name in ("raw-workbook.json", "cambridge-undergraduate.json", "cambridge-undergraduate.csv",
                 "discovery-manifest.json", "excluded-programs.json", "validation-report.json"):
        write("data/cambridge/undergraduate/" + name, (root / DATA / name).read_text(encoding="utf-8"))
    for r in records:
        name = r["id"] + ".json"
        write("data/cambridge/undergraduate/programs/" + name,
              (root / DATA / "programs" / name).read_text(encoding="utf-8"))
    write("reports/cambridge/verification-report.md",
          f"# Cambridge verification report\n\n{report['record_count']} workbook records; "
          f"{report['verified_100_percent_english']} independently certified as fully English-taught; "
          f"{report['language_review_required']} need teaching-language review. "
          "Official Cambridge pages returned HTTP 403. No workbook-reported field is treated as independently verified.\n")
    write("reports/cambridge/duplicate-check.md",
          "# Cambridge duplicate check\n\nNo normalized title-and-qualification duplicates within the 32 workbook records. "
          "The reported three existing Cambridge BTS program names were not provided. "
          "All 32 remain duplicate-check pending; none is marked genuinely new.\n")
    write("reports/cambridge/source-coverage.md",
          f"# Cambridge source coverage\n\n{report['workbook_direct_course_links']} of 32 workbook rows have a direct course URL. "
          "All have Cambridge directory/course-source references. "
          "The official pages returned HTTP 403 in this research session, so field-level official verification remains manual. "
          "The workbook reports 2027-entry fees; the exact tuition academic-year label is not established.\n")
    return expected


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=("ingest", "validate"))
    parser.add_argument("--workbook", type=Path)
    args = parser.parse_args()
    if args.command == "ingest":
        if args.workbook is None:
            parser.error("--workbook is required for ingest")
        print(f"Ingested {len(ingest(args.workbook))} Cambridge workbook rows")
    else:
        print(json.dumps(validate(), indent=2))


if __name__ == "__main__":
    main()
